"""Corrige, por edição direta do XML (sem reescrever o arquivo), as planilhas
"Gestao Financeira <loja> 2026.xlsm":

  1. DRE 2026, linha "LUCROS DISTRIBUIDOS": código 418 na coluna C e, em cada
     coluna mensal com =#REF!, a mesma SUMIFS da linha 37 (401) apontando para $C59.
  2. Cadastros, linha "DIVERSOS - DISTRIBUIÇÃO DE LUCROS": Cod DRE (coluna E) = 418
     quando vazio. (É o PROCV desta tabela que dá o código às linhas de Saídas.)
  3. Cabeçalho da DRE 2026: título com o ano 2026; razão social e CNPJ da loja
     quando faltarem ou estiverem errados (B3/B4), como texto inline.
  4. workbook.xml: fullCalcOnLoad="1" para o Excel recalcular ao abrir.

Só as partes alteradas são regravadas; todo o resto do zip é copiado byte a byte.
Uso local: python corrige_planilha.py ORIGEM DESTINO LOJA   (LOJA = chave em EMPRESAS)
No robô (GitHub Actions) quem chama é corrigir_lote.py.
"""
import html
import logging
import re
import sys
import zipfile
from pathlib import Path

log = logging.getLogger("gerenciais.corrige")

EMPRESAS = {
    "sjp1":   ("3HAT SJP Comércio de Sobremesas Ltda", "19.968.224/0001-10"),
    "ctba3":  ("3HAT Comércio de Sobremesas Ltda", "39.288.388/0001-03"),
    "ctba5":  ("3HAT Jockey Comércio de Sobremesas Ltda", "47.388.413/0001-59"),
    "cl2":    ("3HAT Outlet Comércio de Sobremesas Ltda", "47.958.529/0001-86"),
    "ctba7":  ("3HAT Palladium Comércio de Sobremesas Ltda", "51.906.299/0001-89"),
    "ctba11": ("3HAT Loja Palladium Comércio de Sobremesas Ltda", "61.851.928/0001-74"),
    "mga3":   ("3HT Avenida Comércio de Sobremesas Ltda", "21.605.501/0001-71"),
    "mga5":   ("3HT Cidade Comércio de Sobremesas Ltda", "38.099.619/0001-78"),
    "mga7":   ("3HAT Quiosque Avenida Comércio de Sobremesas Ltda", "58.078.152/0001-97"),
    # razão social continua a do Estação; o rótulo da loja vai junto para ninguém se confundir
    "mga8":   ("3HAT Quiosque Estação Comércio de Sobremesas Ltda — MGA 8 (Havan, ex-CTBA 9)", "55.887.693/0001-78"),
}
# lojas em que B3 é trocado mesmo quando o começo do nome coincide
FORCAR_B3 = {"mga8"}

# trecho do nome do arquivo no Drive → chave da loja. MGA 9, MGA 10 e qualquer
# outro arquivo da pasta ficam de fora até terem linha aqui.
ARQUIVO_PARA_LOJA = {
    "SJP 1": "sjp1", "Ctba 3": "ctba3", "Ctba 5": "ctba5", "Campo Largo 2": "cl2",
    "Ctba 7": "ctba7", "Ctba 11": "ctba11", "Mga 3": "mga3", "Mga 5": "mga5",
    "Mga 7": "mga7", "Mga 8": "mga8",
}


def loja_do_arquivo(nome: str) -> str | None:
    """'Gestao Financeira Mga 3 (Avenida) 2026.xlsm' → 'mga3'; None se não for loja conhecida."""
    if not nome.startswith("Gestao Financeira") or not nome.endswith("2026.xlsm"):
        return None
    for trecho, loja in ARQUIVO_PARA_LOJA.items():
        if f"Gestao Financeira {trecho} " in nome:
            return loja
    return None
MESES = "D F H J L N P R T V X Z".split()


def mapa_abas(z: zipfile.ZipFile) -> dict[str, str]:
    wbx = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid: dict[str, str] = {}
    for m in re.finditer(r"<Relationship\b[^>]*>", rels):
        tag = m.group(0)
        i = re.search(r'Id="([^"]+)"', tag); t = re.search(r'Target="([^"]+)"', tag)
        if i and t:
            rid[i.group(1)] = "xl/" + t.group(1).lstrip("/").removeprefix("xl/")
    return {html.unescape(m.group(1)): rid[m.group(2)]
            for m in re.finditer(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbx)}


def shared_strings(z: zipfile.ZipFile) -> list[str]:
    try:
        x = z.read("xl/sharedStrings.xml").decode("utf-8")
    except KeyError:
        return []
    return [html.unescape(re.sub(r"<[^>]+>", "", si)) for si in re.findall(r"<si>(.*?)</si>", x, re.S)]


def celula(xml: str, ref: str) -> re.Match | None:
    return re.search(rf'<c r="{ref}"(?P<attrs>[^>]*?)(?:/>|>(?P<inner>.*?)</c>)', xml, re.S)


def valor(m: re.Match | None, ss: list[str]) -> str | None:
    if not m:
        return None
    attrs, inner = m.group("attrs"), m.group("inner") or ""
    f = re.search(r"<f[^>]*>(.*?)</f>", inner, re.S)
    if f:
        return "=" + html.unescape(f.group(1))
    v = re.search(r"<v>(.*?)</v>", inner, re.S)
    if v is None:
        t = re.search(r"<t[^>]*>(.*?)</t>", inner, re.S)
        return html.unescape(t.group(1)) if t else None
    if re.search(r't="s"', attrs):
        return ss[int(v.group(1))]
    return html.unescape(v.group(1))


def estilo(m: re.Match | None) -> str:
    s = re.search(r'\bs="(\d+)"', m.group("attrs")) if m else None
    return f' s="{s.group(1)}"' if s else ""


def inline(ref: str, texto: str, s: str) -> str:
    return f'<c r="{ref}"{s} t="inlineStr"><is><t>{html.escape(texto, quote=False)}</t></is></c>'


def substitui_celula(xml: str, ref: str, novo: str) -> str:
    m = celula(xml, ref)
    if m:
        return xml[: m.start()] + novo + xml[m.end():]
    # célula não existe: insere na linha, na posição correta (ordem das colunas)
    col, lin = re.match(r"([A-Z]+)(\d+)", ref).groups()
    row = re.search(rf'<row r="{lin}"[^>]*>(.*?)</row>', xml, re.S)
    if not row:
        raise ValueError(f"linha {lin} não existe para inserir {ref}")
    inner = row.group(1)
    def chave(c: str) -> tuple[int, str]: return (len(c), c)
    pos = len(inner)
    for cm in re.finditer(r'<c r="([A-Z]+)\d+"', inner):
        if chave(cm.group(1)) > chave(col):
            pos = cm.start(); break
    novo_inner = inner[:pos] + novo + inner[pos:]
    return xml[: row.start(1)] + novo_inner + xml[row.end(1):]


SERIAL_2026_INI, SERIAL_2026_FIM = 46023, 46387  # 01/01/2026 e 31/12/2026 no calendário do Excel


def _mes_ano_do_serial(serial: int) -> str:
    """46253 → '8/2026' (texto que a fórmula de Mês Comp. produz)."""
    from datetime import date, timedelta
    d = date(1899, 12, 30) + timedelta(days=serial)
    return f"{d.month}/{d.year}"


def _linhas_2026(xml: str) -> list[tuple[int, int]]:
    """(linha, serial da Data Comp.) de todas as linhas da aba Saídas com Data Comp. em 2026."""
    out: list[tuple[int, int]] = []
    for rm in re.finditer(r'<row r="(\d+)"[^>]*>(.*?)</row>', xml, re.S):
        n = int(rm.group(1))
        a = re.search(rf'<c r="A{n}"[^>]*?(?:/>|>(.*?)</c>)', rm.group(2), re.S)
        v = re.search(r"<v>([\d.]+)</v>", a.group(1)) if a and a.group(1) else None
        if v and SERIAL_2026_INI <= int(float(v.group(1))) <= SERIAL_2026_FIM:
            out.append((n, int(float(v.group(1)))))
    return out


def _linhas_2026_sem_formula_mes(xml: str) -> list[tuple[int, int]]:
    """(linha, serial da Data Comp.) das linhas de 2026 cuja coluna B não é fórmula."""
    out: list[tuple[int, int]] = []
    for rm in re.finditer(r'<row r="(\d+)"[^>]*>(.*?)</row>', xml, re.S):
        n = int(rm.group(1)); inner = rm.group(2)
        a = re.search(rf'<c r="A{n}"[^>]*?(?:/>|>(.*?)</c>)', inner, re.S)
        if not a or not a.group(1):
            continue
        v = re.search(r"<v>([\d.]+)</v>", a.group(1))
        if not v:
            continue
        serial = int(float(v.group(1)))
        if not (SERIAL_2026_INI <= serial <= SERIAL_2026_FIM):
            continue
        b = re.search(rf'<c r="B{n}"[^>]*?(?:/>|>(.*?)</c>)', inner, re.S)
        if b and b.group(1) and "<f" in b.group(1):
            continue
        out.append((n, serial))
    return out


def corrigir_mes_comp(xml: str) -> tuple[str, list[str]]:
    """Coluna 'Mês Comp.' da aba Saídas: linhas de 2026 com texto no lugar da fórmula
    recebem a mesma fórmula explícita das vizinhas. Só mexe nessas células."""
    modelo = re.search(r'<c r="B\d+"[^>]*><f>([^<]+)</f>', xml)
    if not modelo:
        raise ValueError("Saídas: não achei a fórmula de Mês Comp. em nenhuma linha")
    formula = modelo.group(1)  # já escapada como no XML
    mudancas: list[str] = []
    for n, serial in _linhas_2026_sem_formula_mes(xml):
        ref = f"B{n}"
        atual = celula(xml, ref)
        novo = f'<c r="{ref}"{estilo(atual)} t="str"><f>{formula}</f><v>{_mes_ano_do_serial(serial)}</v></c>'
        xml = substitui_celula(xml, ref, novo)
        mudancas.append(f"Saídas {ref}: texto → fórmula de Mês Comp. ({_mes_ano_do_serial(serial)})")
    return xml, mudancas


_MOEDA_BR = re.compile(r"^-?(\d{1,3}(\.\d{3})+|\d+)(,\d{1,2})?$")


def _numero_de_texto_br(texto: str) -> float | None:
    """'\xa04.853,20' / 'R$ 201,00' / 'R 392,27' → float. None se não for valor em reais sem ambiguidade."""
    t = re.sub(r"[\s\xa0]", "", texto)
    t = re.sub(r"^R\$?", "", t)
    if not _MOEDA_BR.match(t):
        return None
    return float(t.replace(".", "").replace(",", "."))


COL_VALOR_SAIDAS = "I"  # Data Comp., Mês Comp., Natureza, Cod DRE, FLCX, Mens., Fornecedor, Descrição, Valor


def _valores_texto_2026(xml: str, ss: list[str]) -> list[tuple[int, str, float | None]]:
    """(linha, texto, número) das linhas de 2026 da aba Saídas cujo Valor está digitado como texto.
    O SUMIFS da DRE ignora texto, então esses lançamentos somem do resultado sem aviso."""
    out: list[tuple[int, str, float | None]] = []
    for n, _serial in _linhas_2026(xml):
        m = celula(xml, f"{COL_VALOR_SAIDAS}{n}")
        if not m or not re.search(r't="(s|inlineStr|str)"', m.group("attrs")):
            continue
        v = valor(m, ss)
        if v is None or v.startswith("=") or not re.search(r"\d", v):
            continue
        out.append((n, v, _numero_de_texto_br(v)))
    return out


def corrigir_valores_texto(xml: str, ss: list[str]) -> tuple[str, list[str]]:
    """Valor digitado como texto vira número (mesmo valor, mesmo estilo). Texto que não é
    valor em reais sem ambiguidade fica como está e sai como ATENÇÃO no relatório."""
    mudancas: list[str] = []
    for n, texto, numero in _valores_texto_2026(xml, ss):
        ref = f"{COL_VALOR_SAIDAS}{n}"
        if numero is None:
            mudancas.append(f"ATENÇÃO Saídas {ref}: valor em texto {texto!r} não reconhecido (não mexi)")
            continue
        xml = substitui_celula(xml, ref, f'<c r="{ref}"{estilo(celula(xml, ref))}><v>{numero:.2f}</v></c>')
        mudancas.append(f"Saídas {ref}: valor em texto {texto!r} → {numero:.2f}")
    return xml, mudancas


# ── Ajustes pontuais aprovados (célula a célula) ─────────────────────────────
# Só aplica se a linha ainda tiver o texto e a data esperados; já corrigida, pula; diferente, avisa.
AJUSTES_DATA_COMP: dict[str, list[dict]] = {
    "ctba5": [
        # André, 12/10/2026 (tracker item 54): comissão de out/2025 lançada com data de 2026.
        {"linha": 7484, "texto": "Reembolso comissão Rose 10/2025", "de": 46326, "para": 45961},
    ],
}


def corrigir_datas_pontuais(xml: str, ss: list[str], loja: str) -> tuple[str, list[str]]:
    mudancas: list[str] = []
    for a in AJUSTES_DATA_COMP.get(loja, []):
        n = a["linha"]
        desc = valor(celula(xml, f"H{n}"), ss) or ""
        atual = valor(celula(xml, f"A{n}"), ss)
        try:
            serial = int(float(atual)) if atual is not None else None
        except ValueError:
            serial = None
        if a["texto"] not in desc:
            mudancas.append(f"ATENÇÃO Saídas A{n}: descrição {desc[:50]!r} não é a esperada; não mexi")
            continue
        if serial == a["para"]:
            continue
        if serial != a["de"]:
            mudancas.append(f"ATENÇÃO Saídas A{n}: data {atual!r} diferente da esperada; não mexi")
            continue
        ref = f"A{n}"
        xml = substitui_celula(xml, ref, f'<c r="{ref}"{estilo(celula(xml, ref))}><v>{a["para"]}</v></c>')
        b = celula(xml, f"B{n}")
        if b and "<f" in (b.group("inner") or ""):
            novo_b = re.sub(r"<v>.*?</v>", f"<v>{_mes_ano_do_serial(a['para'])}</v>", b.group(0), flags=re.S)
            xml = xml[: b.start()] + novo_b + xml[b.end():]
        mudancas.append(
            f"Saídas {ref}: Data Comp. {_mes_ano_do_serial(a['de'])} → {_mes_ano_do_serial(a['para'])} ({a['texto']})"
        )
    return xml, mudancas


def corrigir_receita(dre: str, receita: dict[int, float], ss: list[str]) -> tuple[str, list[str]]:
    """Linha 101 (receita de vendas, linha 7) recebe a receita do BIP360 mês a mês e a
    linha 18 (Receita Bruta Total) volta a ser =SUM(col7:col16) onde estiver digitada."""
    mudancas: list[str] = []
    for mes, esperado in sorted(receita.items()):
        col = MESES[mes - 1]
        ref = f"{col}7"
        atual = valor_num(celula(dre, ref), ss)
        if atual is None or abs(atual - esperado) > 0.5:
            dre = substitui_celula(dre, ref, f'<c r="{ref}"{estilo(celula(dre, ref))}><v>{esperado:.2f}</v></c>')
            mudancas.append(f"DRE {ref} (receita 101, mês {mes}): {atual!r} → {esperado:.2f}")
        ref18 = f"{col}18"
        v18 = valor(celula(dre, ref18), ss)
        if not (isinstance(v18, str) and v18.startswith("=")):
            dre = substitui_celula(dre, ref18, f'<c r="{ref18}"{estilo(celula(dre, ref18))}><f>SUM({col}7:{col}16)</f><v>0</v></c>')
            mudancas.append(f"DRE {ref18} (Receita Bruta Total, mês {mes}): digitado {v18!r} → =SUM({col}7:{col}16)")
    return dre, mudancas


def valor_num(m: re.Match | None, ss: list[str]) -> float | None:
    v = valor(m, ss)
    if v is None or (isinstance(v, str) and v.startswith("=")):
        return None
    try:
        return float(v)
    except ValueError:
        return None


def corrige(origem: str, destino: str, loja: str, receita: dict[int, float] | None = None) -> list[str]:
    z = zipfile.ZipFile(origem)
    abas = mapa_abas(z); ss = shared_strings(z)
    razao, cnpj = EMPRESAS[loja]
    mudancas: list[str] = []
    partes: dict[str, bytes] = {}

    # ── 0. Saídas: Mês Comp. com fórmula em todas as linhas de 2026 ─────────
    p_saidas = abas["Saídas"]
    saidas = z.read(p_saidas).decode("utf-8")
    saidas, m_saidas = corrigir_mes_comp(saidas)
    saidas, m_valores = corrigir_valores_texto(saidas, ss)
    m_saidas += m_valores
    saidas, m_datas = corrigir_datas_pontuais(saidas, ss, loja)
    m_saidas += m_datas
    if any(not m.startswith("ATENÇÃO") for m in m_saidas):
        mudancas.extend(m_saidas)
        partes[p_saidas] = saidas.encode("utf-8")

    # ── 1. DRE 2026 ─────────────────────────────────────────────────────────
    p_dre = abas["DRE 2026"]
    dre = z.read(p_dre).decode("utf-8")
    linha_dist = None
    for n in range(50, 70):
        if "DISTRIB" in str(valor(celula(dre, f"B{n}"), ss) or "").upper() and "SALDO" not in str(valor(celula(dre, f"B{n}"), ss)).upper():
            linha_dist = n; break
    if linha_dist is None:
        raise ValueError("linha LUCROS DISTRIBUIDOS não encontrada na DRE 2026")
    # modelo de fórmula: linha 37 (401), coluna a coluna
    for col in MESES:
        ref = f"{col}{linha_dist}"
        atual = valor(celula(dre, ref), ss)
        modelo = valor(celula(dre, f"{col}37"), ss)
        if not (modelo and modelo.startswith("=SUMIFS")):
            raise ValueError(f"{col}37 não é SUMIFS ({modelo!r}); não sei derivar a fórmula")
        if atual is None or "#REF!" in str(atual) or atual == "":
            formula = modelo[1:].replace("C37", f"$C{linha_dist}")
            if f"$C{linha_dist}" not in formula:
                raise ValueError(f"modelo de {col}37 sem referência a C37: {modelo}")
            novo = f'<c r="{ref}"{estilo(celula(dre, ref))}><f>{html.escape(formula, quote=False)}</f><v>0</v></c>'
            dre = substitui_celula(dre, ref, novo)
            mudancas.append(f"DRE {ref}: {atual!r} → SUMIFS(418)")
    c_ref = f"C{linha_dist}"
    if valor(celula(dre, c_ref), ss) in (None, ""):
        dre = substitui_celula(dre, c_ref, f'<c r="{c_ref}"{estilo(celula(dre, c_ref))}><v>418</v></c>')
        mudancas.append(f"DRE {c_ref}: vazio → 418")
    # ── 3. cabeçalho ────────────────────────────────────────────────────────
    b2 = valor(celula(dre, "B2"), ss) or ""
    if "2026" not in b2:
        dre = substitui_celula(dre, "B2", inline("B2", "DEMONSTRATIVO DO RESULTADO DO EXERCÍCIO - 2026", estilo(celula(dre, "B2"))))
        mudancas.append(f"DRE B2: {b2!r} → 2026")
    b3 = (valor(celula(dre, "B3"), ss) or "").strip()
    def norm(s: str) -> str: return re.sub(r"[^a-z0-9]", "", s.lower().replace("ç", "c").replace("é", "e").replace("ã", "a"))
    if norm(b3)[:14] != norm(razao)[:14] or (loja in FORCAR_B3 and b3 != razao):
        dre = substitui_celula(dre, "B3", inline("B3", razao, estilo(celula(dre, "B3"))))
        mudancas.append(f"DRE B3: {b3!r} → {razao!r}")
    b4 = (valor(celula(dre, "B4"), ss) or "").strip()
    if re.sub(r"\D", "", b4) != re.sub(r"\D", "", cnpj):
        dre = substitui_celula(dre, "B4", inline("B4", cnpj, estilo(celula(dre, "B4"))))
        mudancas.append(f"DRE B4: {b4!r} → {cnpj}")
    # ── 5. receita do BIP360 na linha 101 e Receita Bruta Total como fórmula ──
    if receita:
        dre, m_rec = corrigir_receita(dre, receita, ss)
        mudancas.extend(m_rec)
    partes[p_dre] = dre.encode("utf-8")

    # ── 2. Cadastros ────────────────────────────────────────────────────────
    p_cad = abas.get("Cadastros") or abas.get("Cadastros Iniciais")
    cad = z.read(p_cad).decode("utf-8")
    achou = False
    for rm in re.finditer(r'<row r="(\d+)"[^>]*>(.*?)</row>', cad, re.S):
        n = rm.group(1)
        d = str(valor(celula(cad, f"D{n}"), ss) or "").upper()
        if "DISTRIBUI" in d and "LUCRO" in d:
            achou = True
            e = valor(celula(cad, f"E{n}"), ss)
            if e in (None, ""):
                cad = substitui_celula(cad, f"E{n}", f'<c r="E{n}"{estilo(celula(cad, f"E{n}"))}><v>418</v></c>')
                mudancas.append(f"Cadastros E{n}: vazio → 418")
            elif str(e) not in ("418", "418.0"):
                mudancas.append(f"ATENÇÃO Cadastros E{n} = {e!r} (não mexi)")
            break
    if not achou:
        raise ValueError("linha DIVERSOS - DISTRIBUIÇÃO DE LUCROS não encontrada em Cadastros")
    partes[p_cad] = cad.encode("utf-8")

    # ── 4. recálculo ao abrir ───────────────────────────────────────────────
    wbx = z.read("xl/workbook.xml").decode("utf-8")
    m = re.search(r"<calcPr[^>]*/>", wbx)
    if m and "fullCalcOnLoad" not in m.group(0):
        wbx = wbx.replace(m.group(0), m.group(0)[:-2] + ' fullCalcOnLoad="1"/>')
        partes["xl/workbook.xml"] = wbx.encode("utf-8")

    with zipfile.ZipFile(destino, "w") as out:
        for info in z.infolist():
            out.writestr(info, partes.get(info.filename, z.read(info.filename)), compress_type=info.compress_type)
    return mudancas


def verificar(origem: str, destino: str, loja: str, receita: dict[int, float] | None = None) -> None:
    """Garante que o arquivo corrigido é o original mais as partes esperadas, e que
    as células-alvo ficaram como devem. Levanta ValueError com o motivo."""
    import openpyxl  # só aqui: a correção em si não depende dele

    a, b = zipfile.ZipFile(origem), zipfile.ZipFile(destino)
    if b.testzip() is not None:
        raise ValueError(f"zip corrompido em {destino}: {b.testzip()}")
    if [i.filename for i in a.infolist()] != [i.filename for i in b.infolist()]:
        raise ValueError("lista de partes do zip mudou")
    abas = mapa_abas(a)
    permitidas = {"xl/workbook.xml", abas["DRE 2026"], abas["Saídas"],
                  abas.get("Cadastros") or abas.get("Cadastros Iniciais")}
    sobraram = _linhas_2026_sem_formula_mes(b.read(abas["Saídas"]).decode("utf-8"))
    if sobraram:
        raise ValueError(f"Saídas: {len(sobraram)} linha(s) de 2026 ainda sem fórmula em Mês Comp. (ex.: B{sobraram[0][0]})")
    texto = [(n, t) for n, t, num in _valores_texto_2026(b.read(abas["Saídas"]).decode("utf-8"), shared_strings(b)) if num is not None]
    if texto:
        raise ValueError(f"Saídas: {len(texto)} valor(es) de 2026 ainda em texto (ex.: I{texto[0][0]} = {texto[0][1]!r})")
    saidas_b = b.read(abas["Saídas"]).decode("utf-8")
    ss_b = shared_strings(b)
    for aj in AJUSTES_DATA_COMP.get(loja, []):
        if aj["texto"] in (valor(celula(saidas_b, f"H{aj['linha']}"), ss_b) or ""):
            v = valor(celula(saidas_b, f"A{aj['linha']}"), ss_b)
            if v is None or int(float(v)) != aj["para"]:
                raise ValueError(f"Saídas A{aj['linha']}: Data Comp. não ficou {aj['para']} ({v!r})")
    alteradas = [i.filename for i, j in zip(a.infolist(), b.infolist()) if i.CRC != j.CRC]
    fora = [p for p in alteradas if p not in permitidas]
    if fora:
        raise ValueError(f"partes alteradas fora do esperado: {fora}")
    wb = openpyxl.load_workbook(destino, read_only=True)
    ws = wb["DRE 2026"]
    razao, cnpj = EMPRESAS[loja]
    if "2026" not in str(ws["B2"].value):
        raise ValueError(f"B2 sem 2026: {ws['B2'].value!r}")
    if re.sub(r"\D", "", str(ws["B4"].value or "")) != re.sub(r"\D", "", cnpj):
        raise ValueError(f"B4 ≠ CNPJ da loja: {ws['B4'].value!r}")
    if str(ws["C59"].value) not in ("418", "418.0"):
        raise ValueError(f"C59 ≠ 418: {ws['C59'].value!r}")
    for col in MESES:
        v = str(ws[f"{col}59"].value or "")
        if not v.startswith("=SUMIFS") or "#REF" in v:
            raise ValueError(f"{col}59 não é SUMIFS: {v[:60]!r}")
    for mes, esperado in (receita or {}).items():
        col = MESES[mes - 1]
        v7 = ws[f"{col}7"].value
        if not isinstance(v7, (int, float)) or abs(float(v7) - esperado) > 0.5:
            raise ValueError(f"{col}7 ≠ receita do BIP360 ({esperado:.2f}): {v7!r}")
        v18 = str(ws[f"{col}18"].value or "")
        if not v18.startswith("=SUM("):
            raise ValueError(f"{col}18 não é =SUM: {v18!r}")
    cad = wb["Cadastros"] if "Cadastros" in wb.sheetnames else wb["Cadastros Iniciais"]
    ok = any(
        r[0] and "DISTRIBUI" in str(r[0]).upper() and "LUCRO" in str(r[0]).upper() and str(r[1]) in ("418", "418.0")
        for r in cad.iter_rows(min_row=1, max_row=120, min_col=4, max_col=5, values_only=True)
    )
    if not ok:
        raise ValueError("Cadastros sem 418 em DISTRIBUIÇÃO DE LUCROS")


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
    origem, destino, loja = sys.argv[1:4]
    for linha in corrige(origem, destino, loja):
        log.info("  %s", linha)
    verificar(origem, destino, loja)
    log.info("verificado: %s", Path(destino).name)
